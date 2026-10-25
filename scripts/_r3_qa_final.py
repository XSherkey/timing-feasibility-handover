from pathlib import Path

from docx import Document

D = Path(__file__).resolve().parents[2] / "From Human Intent to In-Orbit Action" / "IEEE_Trans_新版本"
for f in ("Paper_EN_v13.docx", "Paper_ZH_v13.docx"):
    d = Document(D / f)
    txt = "\n".join(p.text for p in d.paragraphs)
    imgs = len([r for r in d.part.rels.values() if "image" in r.reltype])
    print(f, "| paras", len(d.paragraphs), "| tables", len(d.tables), "| images", imgs)
    print("   meta phrases:", txt.count("earlier draft"), txt.count("reviewer"))
    labels = [f"Fig. {n}" for n in range(1, 10)]
    print("   figure labels present:", sum(1 for l in labels if l in txt))
    print("   title:", d.paragraphs[0].text[:80])
