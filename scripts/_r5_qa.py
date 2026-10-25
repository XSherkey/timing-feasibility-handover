"""QA for v7: the split support audit, the corrected percentages, and table integrity."""
from __future__ import annotations

from pathlib import Path

from docx import Document

D = Path(__file__).resolve().parents[2] / "From Human Intent to In-Orbit Action" / "IEEE_Trans_新版本"

for f in ("Paper_EN_v7.docx", "Paper_ZH_v7.docx"):
    doc = Document(D / f)
    paras = [p.text for p in doc.paragraphs]
    text = "\n".join(paras)
    print("=" * 72)
    print(f, "| paras", len(paras), "| tables", len(doc.tables),
          "| images", len([r for r in doc.part.rels.values() if "image" in r.reltype]))
    # split audit tables
    d1a = d1b = None
    for tb in doc.tables:
        head = [c.text.strip() for c in tb.rows[0].cells]
        if len(head) == 6 and ("share" in head[-1].lower() or "比例" in head[-1]):
            first = tb.rows[1].cells[1].text.strip().replace(",", "")
            if first == "3270" or first == "3,270".replace(",", ""):
                d1a = tb
            elif first == "350":
                d1b = tb
    for name, tb in (("D1A", d1a), ("D1B", d1b)):
        if tb is None:
            print("  ", name, "MISSING")
            continue
        tot = 0
        for r in tb.rows[1:-1]:
            tot += int(r.cells[1].text.strip().replace(",", ""))
        last = [c.text.strip() for c in tb.rows[-1].cells]
        print("   %s: rows %d, sum(attempts)=%d, total row=%s" % (name, len(tb.rows), tot, last[:2]))
    # percentage consistency
    checks = {
        "705/8,682": ["705 of 8,682", "705 次"],
        "8.1%": ["8.1%"],
        "no two-orders": ["two orders of magnitude", "两个数量级"],
        "+0.37": ["+0.37"],
        "1.2 pp": ["1.2 pp", "1.2 pp"],
        "eq4-plus": ["min(τ∗ + m, 24 s)"],
    }
    for name, needles in checks.items():
        present = [n for n in needles if n in text]
        print("   %-16s %s" % (name, "present: " + str(len(present)) if present else "absent"))
    print("   '67 of 22,807' occurrences:", text.count("67 of 22,807"), "| '67/22,807':", text.count("67/22,807"))
