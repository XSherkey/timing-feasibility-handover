"""QA for v9: front matter, abstract length, D1 tables, references, code block, figure 2."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document

D = Path(__file__).resolve().parents[2] / "From Human Intent to In-Orbit Action" / "IEEE_Trans_新版本"

for f in ("Paper_EN_v9.docx", "Paper_ZH_v9.docx"):
    doc = Document(D / f)
    paras = [p.text for p in doc.paragraphs]
    text = "\n".join(paras)
    print("=" * 72)
    print(f, "| paras", len(paras), "| tables", len(doc.tables),
          "| images", len([r for r in doc.part.rels.values() if "image" in r.reltype]))
    print("   first paragraph:", paras[0][:80])
    print("   revision note present:", "revision" in text[:400].lower() or "本版" in text[:400])
    # abstract length
    if f.startswith("Paper_EN"):
        i = next(k for k, p in enumerate(paras) if p.strip().lower() == "abstract")
        j = next(k for k, p in enumerate(paras) if p.strip().startswith("Index Terms"))
        words = len(" ".join(paras[i + 1 : j]).split())
        print("   abstract words:", words)
    else:
        i = next(k for k, p in enumerate(paras) if p.strip() == "摘要")
        j = next(k for k, p in enumerate(paras) if p.strip().startswith("关键词"))
        chars = len("".join(paras[i + 1 : j]))
        print("   abstract chars:", chars)
    print("   backslashes in code lines:", sum(1 for p in paras if p.rstrip().endswith("\\")))
    print("   'Table D1A'/'Table D1B':", text.count("Table D1A"), text.count("Table D1B"))
    print("   refs with 'et al.' still:", len([p for p in paras if re.match(r"^\[\d+\]", p.strip()) and "et al." in p]))
