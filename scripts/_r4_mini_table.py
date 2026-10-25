"""Isolate why Table V's body is dropped: build a tiny md with only that block."""
from __future__ import annotations

import subprocess
import sys
from pathlib import Path

D = Path(__file__).resolve().parents[2] / "From Human Intent to In-Orbit Action" / "IEEE_Trans_新版本"
src = (D / "paper_en.md").read_text(encoding="utf-8")
i = src.find("**Table V. Cross-intent")
j = src.find("\n\n", src.find("| transfer gain", i))
block = src[i:j] + "\n"
mini = D / "_mini.md"
mini.write_text("# Mini\n\nIntro paragraph.\n\n" + block, encoding="utf-8")
out = D / "_mini.docx"
node = r"D:\Program Files\Deepseek Harness\resources\runtime\primary-runtime\dependencies\node\bin\node.exe"
r = subprocess.run([node, "build_paper_docx.js", "_mini.md", "_mini.docx", "en"], cwd=D,
                   capture_output=True, text=True, encoding="utf-8", errors="replace")
print(r.stdout[-400:], r.stderr[-400:])

from docx import Document

d = Document(out)
print("tables in mini:", len(d.tables))
for p in d.paragraphs:
    if p.text.strip():
        print("  P:", p.text[:90])
for t in d.tables:
    print("  T:", len(t.rows), "x", len(t.columns),
          [c.text.strip()[:14] for c in t.rows[0].cells])
