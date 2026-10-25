from pathlib import Path

from docx import Document

D = Path(__file__).resolve().parents[2] / "From Human Intent to In-Orbit Action" / "IEEE_Trans_新版本"
for f in ("Paper_EN_v6.docx", "Paper_ZH_v6.docx"):
    doc = Document(D / f)
    print("=" * 70)
    print(f, "tables:", len(doc.tables))
    for i, t in enumerate(doc.tables):
        head = " | ".join(c.text.strip()[:14] for c in t.rows[0].cells)
        print(f"T{i}: {len(t.rows)}x{len(t.columns)}  [{head}]")
        if "bucket" in head.lower() or "桶" in head:
            for r in t.rows:
                print("    ", [c.text.strip() for c in r.cells])
            tot = 0
            for r in t.rows[1:]:
                try:
                    tot += int(r.cells[1].text.strip().replace(",", ""))
                except ValueError:
                    pass
            print("    sum(attempts) =", tot)
