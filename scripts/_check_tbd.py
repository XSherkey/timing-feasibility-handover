import sys, re
from docx import Document
p = sys.argv[1]
d = Document(p)
pat = re.compile(r"\u27e6[^\u27e7]*\u27e7")
def walk(paras, where):
    for i, x in enumerate(paras):
        if pat.search(x.text):
            print(f"[{where} {i}] {x.text[:300]}")
walk(d.paragraphs, "para")
for ti, t in enumerate(d.tables):
    for ri, row in enumerate(t.rows):
        for ci, cell in enumerate(row.cells):
            walk(cell.paragraphs, f"t{ti}r{ri}c{ci}")
