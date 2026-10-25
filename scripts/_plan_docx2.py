"""Rebuild REVISION_PLAN.docx: fix list numbering and join wrapped bullet lines."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

D = Path(__file__).resolve().parents[2] / "From Human Intent to In-Orbit Action" / "IEEE_Trans_新版本"
src = D / "REVISION_PLAN.md"
out = D / "REVISION_PLAN.docx"

doc = Document()
doc.styles["Normal"].font.name = "Times New Roman"
doc.styles["Normal"].font.size = Pt(10.5)

lines = src.read_text(encoding="utf-8").splitlines()
i = 0
num = 0
while i < len(lines):
    raw = lines[i].rstrip()
    if not raw.strip() or raw.strip() == "---":
        i += 1
        continue

    if raw.startswith("# ") or raw.startswith("## ") or raw.startswith("### "):
        level = len(raw) - len(raw.lstrip("#"))
        text = raw.lstrip("#").strip().replace("**", "").replace("*", "")
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt({1: 15, 2: 12.5, 3: 11}[level])
        if level == 1:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            p.paragraph_format.space_before = Pt(10)
        num = 0
        i += 1
        continue

    m = re.match(r"^(\d+)\.\s+(.*)$", raw) or re.match(r"^-\s+(.*)$", raw)
    if m:
        if raw.startswith("-"):
            body = m.group(1)
            prefix = "• "
            num = 0
        else:
            num += 1
            body = m.group(2)
            prefix = f"{num}. "
        # join continuation lines (indented, not a new marker / table / heading)
        i += 1
        while i < len(lines):
            nxt = lines[i]
            if (not nxt.strip() or nxt.startswith("#") or nxt.startswith("|")
                    or re.match(r"^(\d+)\.\s", nxt) or nxt.startswith("- ") or nxt.strip() == "---"):
                break
            body += " " + nxt.strip()
            i += 1
        body = body.replace("**", "").replace("*", "")
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.add_run(prefix + body)
        continue

    if raw.startswith("|"):
        rows = []
        while i < len(lines) and lines[i].startswith("|"):
            cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
            if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                rows.append(cells)
            i += 1
        if rows:
            ncol = max(len(r) for r in rows)
            table = doc.add_table(rows=0, cols=ncol)
            table.style = "Table Grid"
            for ri, row in enumerate(rows):
                cells = table.add_row().cells
                for ci in range(ncol):
                    txt = (row[ci] if ci < len(row) else "").replace("**", "").replace("*", "")
                    cells[ci].text = txt
                    for para in cells[ci].paragraphs:
                        for run in para.runs:
                            run.font.size = Pt(8.5)
                            if ri == 0:
                                run.bold = True
        continue

    txt = raw.replace("**", "").replace("*", "")
    p = doc.add_paragraph(txt)
    p.paragraph_format.space_after = Pt(4)
    i += 1

doc.save(str(out))
print("rebuilt", out, "paragraphs", len(doc.paragraphs), "tables", len(doc.tables))
