from docx import Document
p = r"D:\Desktop_Files\论文\1. LLM Satellite Reinforcement Learning\自己写\From Human Intent to In-Orbit Action\EAAI\Revised_Manuscript_实测数据填充版.docx"
d = Document(p)
keys = ["deepseek-flash", "59.9", "88.9", "0.045 ms", "29.2", "0.665 dB", "50.6", "1.417"]
text = "\n".join(x.text for x in d.paragraphs)
for t in d.tables:
    for row in t.rows:
        for c in row.cells:
            text += "\n" + c.text
print("markers left:", text.count("\u27e6"))
for k in keys:
    print(f"  contains {k!r}: {k in text}")
# show the first filled sentence
for x in d.paragraphs:
    if "deepseek-flash" in x.text:
        print("\nfirst mention:\n", x.text[:400])
        break
