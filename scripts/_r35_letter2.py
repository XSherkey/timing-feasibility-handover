"""Render the cover letter with a leading bold phrase per paragraph."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

D = Path(__file__).resolve().parents[2] / "From Human Intent to In-Orbit Action" / "IEEE_Trans_新版本"
src = D / "COVER_LETTER.md"
out = D / "COVER_LETTER.docx"

doc = Document()
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Inches(0.9)
sec.top_margin = sec.bottom_margin = Inches(0.8)
normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(10)
normal.paragraph_format.space_after = Pt(5)

lines = src.read_text(encoding="utf-8").splitlines()
i = 0
while i < len(lines):
    raw = lines[i].rstrip()
    if not raw.strip() or raw.strip() == "---":
        i += 1
        continue
    if raw.startswith("# "):
        p = doc.add_paragraph()
        r = p.add_run(raw[2:].strip())
        r.bold = True
        r.font.size = Pt(11.5)
        p.paragraph_format.space_after = Pt(6)
        i += 1
        continue
    if raw.startswith("**") and raw.count("**") == 2 and raw.rstrip().endswith("**"):
        p = doc.add_paragraph()
        r = p.add_run(raw.strip().strip("*"))
        r.font.size = Pt(10)
        p.paragraph_format.space_after = Pt(2)
        i += 1
        continue
    buf = [raw]
    i += 1
    while i < len(lines) and lines[i].strip() and not lines[i].startswith(("#", "**", "---")):
        buf.append(lines[i].strip())
        i += 1
    text = " ".join(buf)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    m = re.match(r"^\*\*(.+?)\.\*\*\s*(.*)$", text)
    if m:
        r1 = p.add_run(m.group(1) + ".")
        r1.bold = True
        import html as _html
        rest = _html.unescape(m.group(2).replace("**", "").replace("*", ""))
        if rest:
            p.add_run(" " + rest)
    else:
        try:
            import html
            text = html.unescape(text)
        except Exception:
            pass
        p.add_run(text.replace("**", "").replace("*", ""))

doc.save(str(out))
print("written", out, "paragraphs", len(doc.paragraphs))
