"""QA for the v6 build: the fourth review's items, table integrity, and number consistency."""
from __future__ import annotations

from pathlib import Path

from docx import Document

D = Path(__file__).resolve().parents[2] / "From Human Intent to In-Orbit Action" / "IEEE_Trans_新版本"

CHECKS = {
    "D1-counts": ["92.5% sit at or above 16 s", "92.5% 的尝试位于 16 s 及以上"],
    "near-polar-band": ["near-polar band that the rule scopes", "近极区带"],
    "VIA-no-contradiction": ["The lockbox comparison is executed in the next paragraph",
                             "该 lockbox 比较已在下一段执行"],
    "lim7-bounded": ["is bounded", "收窄"],
    "probe-5seed": ["five policy seeds at each of three physics regimes",
                    "五个策略种子、三档物理"],
    "pairing-split": ["the apples-to-apples pairing", "同口径配对"],
    "contrib2-KM": ["serves as a diagnostic consistency check, not as the source",
                    "只作为诊断性一致性检查"],
    "intro-effective": ["effective-support failure under a sample budget",
                        "给定样本预算下的有效支撑失败"],
    "ledger-in-table2": ["mined-rule campaign arm (16 s, no margin)", "挖出规则 campaign 臂"],
}

for f in ("Paper_EN_v6.docx", "Paper_ZH_v6.docx"):
    doc = Document(D / f)
    text = "\n".join(p.text for p in doc.paragraphs)
    table_text = []
    for tb in doc.tables:
        for r in tb.rows:
            for c in r.cells:
                table_text.append(c.text)
    blob = (text + "\n" + "\n".join(table_text)).lower()
    print("==", f, "| paras", len(doc.paragraphs), "| tables", len(doc.tables),
          "| images", len([r for r in doc.part.rels.values() if "image" in r.reltype]))
    for name, needles in CHECKS.items():
        ok = any(n.lower() in blob for n in needles)
        print("   %-24s %s" % (name, "OK" if ok else "MISSING"))
    print("   literal '**' paragraphs:", sum(1 for p in doc.paragraphs if "**" in p.text))
    # table D1 sanity: find the 8x6 table and check the 16-18 row
    for tb in doc.tables:
        if len(tb.columns) == 6 and len(tb.rows) == 8:
            head = [c.text.strip() for c in tb.rows[0].cells]
            if "attempts" in head[1].lower() or "尝试" in head[1]:
                print("   D1 rows 14-16 / 16-18:", [[c.text.strip() for c in tb.rows[k].cells] for k in (5, 6)])
