from pathlib import Path

from docx import Document

D = Path(__file__).resolve().parents[2] / "From Human Intent to In-Orbit Action" / "IEEE_Trans_新版本"
for f in ("Paper_EN_v5.docx", "Paper_ZH_v5.docx"):
    d = Document(D / f)
    print("==", f, "| paragraphs", len(d.paragraphs), "| tables", len(d.tables),
          "| images", len([r for r in d.part.rels.values() if "image" in r.reltype]))
    for i, t in enumerate(d.tables):
        head = " | ".join(c.text.strip()[:16] for c in t.rows[0].cells)
        widths = {len(r.cells) for r in t.rows}
        flag = "" if len(widths) == 1 else "  <-- RAGGED"
        print("  T%d: %dx%d head=[%s]%s" % (i, len(t.rows), len(t.columns), head[:80], flag))
