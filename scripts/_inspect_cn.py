from docx import Document
p = r"D:\Desktop_Files\论文\1. LLM Satellite Reinforcement Learning\自己写\From Human Intent to In-Orbit Action\EAAI\复现结果\复现报告_中文.docx"
d = Document(p)
for i, para in enumerate(d.paragraphs[:8]):
    print(i, repr(para.text[:120]))
print("--- any asterisks in body? ---")
n = sum(1 for para in d.paragraphs if "**" in para.text)
print("paragraphs containing '**':", n)
print("paragraphs containing '---':", sum(1 for para in d.paragraphs if para.text.strip() == "---"))
