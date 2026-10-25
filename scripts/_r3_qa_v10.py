from pathlib import Path

from docx import Document

D = Path(__file__).resolve().parents[2] / "From Human Intent to In-Orbit Action" / "IEEE_Trans_新版本"
for f in ("Paper_EN_v11.docx", "Paper_ZH_v11.docx"):
    d = Document(D / f)
    txt = "\n".join(p.text for p in d.paragraphs)
    imgs = len([r for r in d.part.rels.values() if "image" in r.reltype])
    print(f, "| paras", len(d.paragraphs), "| tables", len(d.tables), "| images", imgs)
    print("   meta phrases:", txt.count("earlier draft"), txt.count("earlier version"), txt.count("reviewer"))
    print("   Table VII/VIII:", "Table VII." in txt, "Table VIII." in txt, "| 表 VII:", "表 VII." in txt)
    print("   nomenclature:", "Nomenclature" in txt or "符号表" in txt)
    print("   title:", d.paragraphs[0].text[:90])
