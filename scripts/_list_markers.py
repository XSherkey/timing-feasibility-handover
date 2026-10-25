import re, sys
from docx import Document
p = r"D:\Desktop_Files\论文\1. LLM Satellite Reinforcement Learning\自己写\From Human Intent to In-Orbit Action\EAAI\Revised_Manuscript_修改稿.docx"
d = Document(p)
pat = re.compile(r"\u27e6[^\u27e7]*\u27e7")
n = 0
for i, x in enumerate(d.paragraphs):
    ms = pat.findall(x.text)
    if ms:
        n += len(ms)
        print(f"[{i}] {len(ms)} marker(s): {x.text[:150]} ...")
for ti, t in enumerate(d.tables):
    for ri, row in enumerate(t.rows):
        for ci, c in enumerate(row.cells):
            for x in c.paragraphs:
                ms = pat.findall(x.text)
                if ms:
                    n += len(ms)
                    print(f"[t{ti}r{ri}c{ci}] {len(ms)}: {x.text[:120]}")
print("total markers:", n)
