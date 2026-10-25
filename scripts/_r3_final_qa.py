"""Final QA for the v5 build: check the seven review items are present, that no
stray markdown emphasis leaked into the docx, and that the tables are intact."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document

D = Path(__file__).resolve().parents[2] / "From Human Intent to In-Orbit Action" / "IEEE_Trans_新版本"

CHECKS = {
    "P3-inclusion": ["{T_z > max(a, τ)} ⊆ {T_z > τ}", "包含关系"],
    "P2-effective-support": ["effective-support failure under a sample budget", "给定样本预算下的有效支撑失败"],
    "P1-plateau18": ["the first grid point at which the measured late-handover share is exactly zero is ā = 18 s",
                     "实测晚切换占比首次精确为 0 的网格点"],
    "ledger-names": ["mined-rule campaign arm", "挖出规则 campaign 臂"],
    "lockbox-tolerance": ["the tolerance is **1.2 pp**", "容差为 **1.2 pp**"],
    "gate-count": ["four of the six degrade held-out success outright", "四条直接降低留出集成功率"],
    "positioning": ["systems-and-audit contribution with formal scope notes",
                    "带有明确适用范围注记的系统与审计"],
}

for f, lang in (("Paper_EN_v5.docx", 0), ("Paper_ZH_v5.docx", 1)):
    doc = Document(D / f)
    text = "\n".join(p.text for p in doc.paragraphs)
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                text += "\n" + c.text
    print("==", f, "| chars", len(text), "| tables", len(doc.tables),
          "| images", len([r for r in doc.part.rels.values() if "image" in r.reltype]))
    for name, needles in CHECKS.items():
        ok = any(n in text for n in needles)
        print("   %-22s %s" % (name, "OK" if ok else "MISSING"))
    stray = re.findall(r"\S{0,12}[*]\S{0,12}", text)
    print("   stray asterisks:", len(stray), stray[:5])
