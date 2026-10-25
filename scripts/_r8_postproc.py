"""Post-process v2: keep every table on one page and glue captions to their tables.

This is what the last three reviews were really complaining about: a table that breaks across a page makes
any PDF-to-text extractor fuse the last cell of one fragment with the first cell of the next, which then
looks like "3,508" or "1,096" and like invented captions ("Table 1.1.1...", "Table 1: hours of sleep").
"""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

D = Path(__file__).resolve().parents[2] / "From Human Intent to In-Orbit Action" / "IEEE_Trans_新版本"


def keep(p: Paragraph) -> None:
    p.paragraph_format.keep_with_next = True


for name in sys.argv[1:]:
    path = D / name
    doc = Document(path)

    # 1. every row keeps with the next one, so the table cannot split across a page
    for tb in doc.tables:
        rows = tb.rows
        for i, row in enumerate(rows):
            trPr = row._tr.get_or_add_trPr()
            if trPr.find(qn("w:cantSplit")) is None:
                trPr.append(OxmlElement("w:cantSplit"))
            if i == 0 and trPr.find(qn("w:tblHeader")) is None:
                trPr.append(OxmlElement("w:tblHeader"))
            if i < len(rows) - 1:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        keep(p)

    # 2. a caption paragraph keeps with what follows it
    body = doc.element.body
    for child in body.iterchildren():
        if child.tag.endswith("}p"):
            p = Paragraph(child, doc)
            txt = p.text.strip()
            if txt.startswith(("Table ", "Fig. ", "表 ", "图 ")):
                keep(p)

    doc.save(path)
    print("post-processed", name, "| tables", len(doc.tables))
