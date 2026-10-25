"""Post-process the built docx: keep table rows from splitting across pages and repeat header rows."""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

D = Path(__file__).resolve().parents[2] / "From Human Intent to In-Orbit Action" / "IEEE_Trans_新版本"

for name in sys.argv[1:]:
    path = D / name
    doc = Document(path)
    for tb in doc.tables:
        for r_i, row in enumerate(tb.rows):
            trPr = row._tr.get_or_add_trPr()
            if trPr.find(qn("w:cantSplit")) is None:
                trPr.append(OxmlElement("w:cantSplit"))
            if r_i == 0 and trPr.find(qn("w:tblHeader")) is None:
                trPr.append(OxmlElement("w:tblHeader"))
    doc.save(path)
    print("post-processed", name, "tables:", len(doc.tables))
