"""Bisect: find the first prefix of the paper whose build contains Table V."""
from __future__ import annotations

import subprocess
from pathlib import Path

from docx import Document

D = Path(__file__).resolve().parents[2] / "From Human Intent to In-Orbit Action" / "IEEE_Trans_新版本"
node = r"D:\Program Files\Deepseek Harness\resources\runtime\primary-runtime\dependencies\node\bin\node.exe"
src = (D / "paper_en.md").read_text(encoding="utf-8")
cap = src.find("**Table V. Cross-intent")
tbl_end = src.find("\n\n", src.find("| transfer gain", cap))


def build(prefix: str, tag: str) -> bool:
    mini = D / f"_bis_{tag}.md"
    out = D / f"_bis_{tag}.docx"
    mini.write_text(prefix + "\n## References\n\n[end]\n", encoding="utf-8")
    subprocess.run([node, "build_paper_docx.js", mini.name, out.name, "en"], cwd=D,
                   capture_output=True, text=True, encoding="utf-8", errors="replace")
    d = Document(out)
    ok = any("intent 1 latency" in c.text or "transfer gain" in c.text
             for t in d.tables for r in t.rows for c in r.cells)
    print(f"{tag}: tables={len(d.tables)} tableV={ok}")
    return ok


build(src[: tbl_end], "full_prefix")
build(src[:cap] + src[cap:tbl_end], "upto_table")
# halve the prefix until it works
lo, hi = 0, cap
for k in range(6):
    mid = (lo + hi) // 2
    ok = build(src[:mid] + "\n\n" + src[cap:tbl_end], f"mid{mid}")
    if ok:
        hi = mid
    else:
        lo = mid
print("first failing boundary between", lo, "and", hi)
