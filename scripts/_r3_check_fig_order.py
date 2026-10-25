"""Final check: figure citation order must be monotone, and no figure may follow the references."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document

D = Path(__file__).resolve().parents[2] / "From Human Intent to In-Orbit Action" / "IEEE_Trans_新版本"
for name in ("paper_en.md", "paper_zh.md"):
    t = (D / name).read_text(encoding="utf-8")
    cut = t.find("## References" if name.endswith("en.md") else "## 参考文献")
    body = t[:cut]
    pat = r"Fig\.\s*(\d+)" if name.endswith("en.md") else r"图\s*(\d+)"
    order = []
    for m in re.finditer(pat, body):
        n = int(m.group(1))
        if n not in order:
            order.append(n)
    print(name, "| first-mention order:", order,
          "| monotone:", order == sorted(order))

for f in ("Paper_EN_v14.docx", "Paper_ZH_v14.docx"):
    d = Document(D / f)
    paras = d.paragraphs
    refs_idx = next((i for i, p in enumerate(paras)
                     if p.text.strip().startswith(("References", "参考文献"))), None)
    after = 0
    if refs_idx is not None:
        body_el = d.element.body
        seq = list(body_el.iterchildren())
        # count images after the references paragraph
        seen_refs = False
        for child in seq:
            txt = "".join(child.itertext())
            if not seen_refs and txt.strip().startswith(("References", "参考文献")):
                seen_refs = True
                continue
            if seen_refs and child.findall(".//{http://schemas.openxmlformats.org/drawingml/2006/main}blip"):
                after += 1
    imgs = len([r for r in d.part.rels.values() if "image" in r.reltype])
    print(f, "| images", imgs, "| images after references:", after,
          "| tables", len(d.tables))
