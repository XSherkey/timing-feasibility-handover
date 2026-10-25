"""One-process finalisation of the document deliverables.

Does everything that needs ``python-docx`` in a single interpreter start, because
importing the scientific stack on this machine costs minutes:

1. refills the manuscript's TBD markers from ``docs/tbd_values.json``;
2. converts the Chinese v3 report to .docx with the author's own converter;
3. prints a verification summary (markers left, key values present).
"""
from __future__ import annotations

import importlib.util
import json
import re
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
EAAI = Path(
    r"D:\Desktop_Files\论文\1. LLM Satellite Reinforcement Learning\自己写"
    r"\From Human Intent to In-Orbit Action\EAAI"
)

sys.path.insert(0, str(ROOT / "scripts"))
from tbd_rules import build_rules  # noqa: E402


def load_module(path: Path, name: str):
    spec = importlib.util.spec_from_file_location(name, path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def clean_md(text: str) -> str:
    out = []
    for line in text.split("\n"):
        if re.fullmatch(r"\s*-{3,}\s*", line):
            out.append("")
            continue
        line = re.sub(r"(?<!\*)\*([^*\n]+)\*(?!\*)", r"\1", line)
        out.append(line)
    return "\n".join(out)


def main() -> int:
    tbd = json.loads((ROOT / "docs/tbd_values.json").read_text(encoding="utf-8"))
    v = lambda d, k, default="n/a": (d.get(k) or {}).get("value", default)  # noqa: E731

    # the convergence rules need ep90 medians, which finalize2 does not compute
    camp = ROOT / "results" / "campaign"
    per: dict[str, list[float]] = {}
    for p in camp.glob("*/run_summary.json"):
        try:
            with open(p, encoding="utf-8") as fh:
                r = json.load(fh)
        except (json.JSONDecodeError, OSError):
            continue
        if not r.get("ok"):
            continue
        curve = r.get("curve") or []
        if len(curve) < 2:
            continue
        ys = [c.get("S", float("nan")) for c in curve]
        xs = [c.get("episodes", 0) for c in curve]
        finite = [y for y in ys if y == y]
        if not finite:
            continue
        target = 0.9 * (sum(finite[-2:]) / len(finite[-2:]))
        hit = [x for x, y in zip(xs, ys) if y == y and y >= target]
        per.setdefault(r["method"], []).append(float(hit[0]) if hit else float(xs[-1]))
    for m, vals in per.items():
        vals = sorted(vals)
        med = vals[len(vals) // 2]
        tbd[f"ep90.{m}"] = dict(value=f"{med:.0f} episodes (median over {len(vals)} seeds)",
                                source=f"campaign/{m}")

    rules = build_rules(tbd, v)

    fill = load_module(ROOT / "scripts" / "fill_docx.py", "fill_docx_mod")
    src = EAAI / "Revised_Manuscript_修改稿.docx"
    dst = EAAI / "Revised_Manuscript_实测数据填充版.docx"
    stats = fill.fill_doc(src, dst, tbd)
    print("manuscript fill:", stats)

    conv = load_module(EAAI / "md2docx_local.py", "md2docx").convert
    md = ROOT / "docs" / "REPRODUCTION_REPORT_中文_v3.md"
    cleaned = clean_md(md.read_text(encoding="utf-8"))
    tmp = Path(tempfile.gettempdir()) / "almond_v3_cn.md"
    tmp.write_text(cleaned, encoding="utf-8")
    out_docx = EAAI / "复现结果" / "复现报告_中文_v3.docx"
    out_docx.parent.mkdir(parents=True, exist_ok=True)
    conv(str(tmp), str(out_docx), title_east="SimHei")
    print("chinese report:", out_docx.name, out_docx.stat().st_size, "bytes")

    from docx import Document

    d = Document(str(dst))
    text = "\n".join(p.text for p in d.paragraphs)
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                text += "\n" + c.text
    print("markers left in manuscript:", text.count("\u27e6"))
    for k in ("93.6", "72.6", "65.0", "1.159", "80%", "deepseek-flash", "0.045"):
        print(f"  contains {k!r}: {k in text}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
