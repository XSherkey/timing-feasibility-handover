"""Turn REVISION_PLAN.md into a readable .docx for the reviewer (headings, bullets, tables)."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

D = Path(__file__).resolve().parents[2] / "From Human Intent to In-Orbit Action" / "IEEE_Trans_新版本"
src = D / "REVISION_PLAN.md"
out = D / "REVISION_PLAN.docx"

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(10.5)

lines = src.read_text(encoding="utf-8").splitlines()
i = 0
while i < len(lines):
    line = lines[i].rstrip()
    if not line.strip():
        i += 1
        continue
    if line.strip() == "---":
        i += 1
        continue
    if line.startswith("# "):
        p = doc.add_paragraph()
        r = p.add_run(line[2:].strip())
        r.bold = True
        r.font.size = Pt(15)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        i += 1
        continue
    if line.startswith("## "):
        p = doc.add_paragraph()
        r = p.add_run(line[3:].strip())
        r.bold = True
        r.font.size = Pt(12.5)
        p.paragraph_format.space_before = Pt(10)
        i += 1
        continue
    if line.startswith("### "):
        p = doc.add_paragraph()
        r = p.add_run(line[4:].strip())
        r.bold = True
        r.font.size = Pt(11)
        p.paragraph_format.space_before = Pt(8)
        i += 1
        continue
    if line.startswith("|"):
        rows = []
        while i < len(lines) and lines[i].startswith("|"):
            cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
            if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                rows.append(cells)
            i += 1
        if rows:
            ncol = max(len(r) for r in rows)
            table = doc.add_table(rows=0, cols=ncol)
            table.style = "Table Grid"
            # column widths proportional to the longest cell in each column
            lens = [max(len(r[c]) if c < len(r) else 0 for r in rows) for c in range(ncol)]
            total = sum(lens) or ncol
            widths = [max(0.8, 6.9 * l / total) for l in lens]
            table.autofit = False
            for ri, row in enumerate(rows):
                cells = table.add_row().cells
                for ci in range(ncol):
                    txt = row[ci] if ci < len(row) else ""
                    txt = txt.replace("**", "").replace("*", "")
                    cells[ci].text = txt
                    cells[ci].width = __import__("docx").shared.Inches(widths[ci])
                    for para in cells[ci].paragraphs:
                        for run in para.runs:
                            run.font.size = Pt(8.5)
                            if ri == 0:
                                run.bold = True
        continue
    if line.startswith("- ") or re.match(r"^\d+\.\s", line):
        txt = re.sub(r"^(- |\d+\.\s)", "", line)
        txt = txt.replace("**", "").replace("*", "")
        p = doc.add_paragraph(txt, style="List Bullet" if line.startswith("- ") else "List Number")
        for run in p.runs:
            run.font.size = Pt(10.5)
        i += 1
        continue
    txt = line.replace("**", "").replace("*", "")
    p = doc.add_paragraph(txt)
    p.paragraph_format.space_after = Pt(4)
    i += 1

doc.save(str(out))
print("written", out, "paragraphs", len(doc.paragraphs), "tables", len(doc.tables))
