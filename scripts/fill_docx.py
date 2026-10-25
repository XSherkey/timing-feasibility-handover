"""Fill the manuscript's TBD markers with the measured values.

Reads ``docs/tbd_values.json`` (produced by ``make_report.py``) and writes a new
DOCX next to the original, replacing every ``⟦TBD...⟧`` marker with the measured
value together with its provenance.  The original file is never modified.

Rules are matched on a distinctive phrase of the surrounding paragraph (or table
cell), because several markers share the same generic ``⟦TBD-DATA⟧`` spelling.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
MARKER = re.compile(r"⟦[^⟧]*⟧")


def load_tbd(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def v(tbd: dict, key: str, default: str = "n/a") -> str:
    e = tbd.get(key)
    return e["value"] if e else default


def build_rules(tbd: dict) -> list[tuple[str, list[str]]]:
    """Delegate to scripts/tbd_rules.py so the mapping stays reviewable in one place."""
    import sys
    from pathlib import Path as _P

    sys.path.insert(0, str(_P(__file__).resolve().parent))
    from tbd_rules import build_rules as _build

    return _build(tbd, lambda d, k, default="n/a": (d.get(k) or {}).get("value", default))


def fill_doc(src: Path, dst: Path, tbd: dict) -> dict:
    doc = Document(str(src))
    rules = build_rules(tbd)
    stats = dict(markers=0, replaced=0, unresolved=[])

    def process(paragraph, row_context: str = "") -> None:
        text = paragraph.text
        markers = MARKER.findall(text)
        if not markers:
            return
        stats["markers"] += len(markers)
        values = None
        haystack = text + " || " + row_context
        for ctx, vals in rules:
            if ctx in haystack:
                values = vals
                break
        if values is None:
            for m in markers:
                stats["unresolved"].append(m)
            return
        used = 0
        for m in markers:
            val = values[used] if used < len(values) else values[-1]
            if m in text and val:
                text = text.replace(m, val, 1)
                used += 1
                stats["replaced"] += 1
        # rewrite the paragraph while keeping its style
        if paragraph.runs:
            paragraph.runs[0].text = text
            for r in paragraph.runs[1:]:
                r.text = ""
        else:
            paragraph.text = text

    for p in doc.paragraphs:
        process(p)
    for tbl in doc.tables:
        for row in tbl.rows:
            row_ctx = " ".join(c.text for c in row.cells)
            for cell in row.cells:
                for p in cell.paragraphs:
                    process(p, row_ctx)
    for section in doc.sections:
        for p in section.header.paragraphs:
            process(p)
        for p in section.footer.paragraphs:
            process(p)
    doc.save(str(dst))
    return stats


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", required=True)
    ap.add_argument("--tbd", default="docs/tbd_values.json")
    ap.add_argument("--dst", default=None)
    args = ap.parse_args()
    src = Path(args.src)
    dst = Path(args.dst) if args.dst else src.with_name(src.stem + "_实测数据填充版.docx")
    tbd = load_tbd(ROOT / args.tbd)
    stats = fill_doc(src, dst, tbd)
    print(json.dumps(stats, indent=2, ensure_ascii=False))
    print("wrote", dst)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
